"""
generate_client_resumes.py - Generates professionally formatted DOCX resume templates
for registered client candidates: Safi Ahmed, Tanuja Sharma, and Jane Doe.
"""

from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from client_config import CLIENTS_DIR


def set_run_font(run, font_name="Calibri", font_size=11, bold=False, italic=False, color_rgb=(51, 51, 51)):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)


def create_header(doc, name, contact_line):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p.add_run(f"{name}\n")
    set_run_font(r_name, font_name="Calibri", font_size=20, bold=True, color_rgb=(24, 43, 73))
    
    r_contact = p.add_run(contact_line)
    set_run_font(r_contact, font_name="Calibri", font_size=10, color_rgb=(90, 90, 90))


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    set_run_font(run, font_name="Calibri", font_size=12, bold=True, color_rgb=(24, 76, 120))


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, font_name="Calibri", font_size=10.5, color_rgb=(51, 51, 51))
    return p


def build_safi_resume():
    doc = Document()
    create_header(
        doc,
        name="SAFI AHMED",
        contact_line="Dallas, TX | +1 (469) 555-0142 | safi.ahmed.c2c@example.com | linkedin.com/in/safi-ahmed-dev | Work Auth: C2C"
    )

    add_section_heading(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    r_sum = p_sum.add_run(
        "Accomplished Senior Java Full Stack Developer with 8+ years of extensive hands-on experience designing, "
        "developing, and deploying enterprise-grade distributed systems and microservices. Proven expertise in Core Java, "
        "Spring Boot, REST APIs, Kafka event-driven architectures, PostgreSQL, and AWS cloud infrastructures."
    )
    set_run_font(r_sum, font_size=10.5)

    add_section_heading(doc, "Technical Skills")
    table = doc.add_table(rows=4, cols=2)
    skills = [
        ("Languages & Frameworks:", "Java 17/11/8, Spring Boot, Spring Cloud, Hibernate, Microservices, Python, SQL"),
        ("Frontend Technologies:", "React.js, TypeScript, JavaScript (ES6+), HTML5, CSS3, Redux Toolkit"),
        ("Databases & Messaging:", "PostgreSQL, MySQL, Oracle, MongoDB, Redis, Apache Kafka, RabbitMQ"),
        ("Cloud & DevOps:", "AWS (EC2, S3, RDS, ECS), Docker, Kubernetes, Jenkins, CI/CD, Terraform, Git")
    ]
    for i, (cat, val) in enumerate(skills):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        r0 = c0.paragraphs[0].add_run(cat)
        set_run_font(r0, bold=True, font_size=10, color_rgb=(24, 43, 73))
        r1 = c1.paragraphs[0].add_run(val)
        set_run_font(r1, font_size=10)

    add_section_heading(doc, "Professional Experience")
    
    # Role 1
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(8)
    r1 = p1.add_run("Lead Java Full Stack Developer | Enterprise Health Solutions, Dallas, TX (2021 – Present)")
    set_run_font(r1, bold=True, font_size=11, color_rgb=(24, 43, 73))
    add_bullet(doc, "Architected high-throughput microservices using Java 17, Spring Boot, and PostgreSQL handling 15M+ daily transactions.")
    add_bullet(doc, "Engineered event-driven decoupled asynchronous messaging pipelines with Apache Kafka and RabbitMQ.")
    add_bullet(doc, "Containerized microservice workloads using Docker and orchestrated deployments on AWS ECS and Kubernetes clusters.")
    add_bullet(doc, "Built responsive, accessible management portals with React, TypeScript, and modern Redux state management.")

    # Role 2
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    r2 = p2.add_run("Senior Java Developer | Global Financial Tech, New York, NY (2018 – 2021)")
    set_run_font(r2, bold=True, font_size=11, color_rgb=(24, 43, 73))
    add_bullet(doc, "Developed scalable RESTful APIs utilizing Spring MVC, Spring Data JPA, and Hibernate with PostgreSQL backends.")
    add_bullet(doc, "Decreased query latency by 45% through Redis caching strategies and optimized database indexing.")
    add_bullet(doc, "Maintained automated CI/CD deployment pipelines using Jenkins, Git, and Docker.")

    add_section_heading(doc, "Education")
    p_edu = doc.add_paragraph()
    r_edu = p_edu.add_run("Bachelor of Science in Computer Science | University of Texas at Dallas")
    set_run_font(r_edu, font_size=10)

    out_file = CLIENTS_DIR / "safi_resume.docx"
    doc.save(out_file)
    print(f"Created Safi resume: {out_file}")


def build_tanuja_resume():
    doc = Document()
    create_header(
        doc,
        name="TANUJA SHARMA",
        contact_line="Austin, TX | +1 (512) 555-0188 | tanuja.sharma.ml@example.com | linkedin.com/in/tanuja-sharma-ai | Work Auth: C2C"
    )

    add_section_heading(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    r_sum = p_sum.add_run(
        "Senior Data Scientist and Machine Learning Engineer with 7+ years of experience leading the development and "
        "deployment of production AI/ML systems, NLP pipelines, and Large Language Model (LLM) architectures. "
        "Skilled in Python, PyTorch, Scikit-Learn, Databricks, Snowflake, and AWS MLOps."
    )
    set_run_font(r_sum, font_size=10.5)

    add_section_heading(doc, "Technical Skills")
    table = doc.add_table(rows=4, cols=2)
    skills = [
        ("Machine Learning & AI:", "PyTorch, TensorFlow, Scikit-Learn, Hugging Face, LLMs, NLP, LangChain, Computer Vision"),
        ("Languages & Libraries:", "Python, SQL, R, Pandas, NumPy, SciPy, Matplotlib, Seaborn, FastAPI"),
        ("Big Data & Storage:", "Databricks, Apache Spark, Snowflake, PostgreSQL, BigQuery, AWS S3"),
        ("MLOps & Cloud:", "AWS SageMaker, Docker, Kubernetes, MLflow, Airflow, CI/CD, Git")
    ]
    for i, (cat, val) in enumerate(skills):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        r0 = c0.paragraphs[0].add_run(cat)
        set_run_font(r0, bold=True, font_size=10, color_rgb=(24, 43, 73))
        r1 = c1.paragraphs[0].add_run(val)
        set_run_font(r1, font_size=10)

    add_section_heading(doc, "Professional Experience")
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(8)
    r1 = p1.add_run("Lead Machine Learning Engineer | DataVision Analytics, Austin, TX (2022 – Present)")
    set_run_font(r1, bold=True, font_size=11, color_rgb=(24, 43, 73))
    add_bullet(doc, "Designed and deployed generative AI models and retrieval-augmented generation (RAG) pipelines using Python and LangChain.")
    add_bullet(doc, "Built scalable machine learning workflows using PyTorch and Scikit-Learn on Databricks clusters.")
    add_bullet(doc, "Automated end-to-end model training, evaluation, and registry deployment using MLflow and AWS SageMaker.")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    r2 = p2.add_run("Senior Data Scientist | FinEdge Solutions, San Jose, CA (2019 – 2022)")
    set_run_font(r2, bold=True, font_size=11, color_rgb=(24, 43, 73))
    add_bullet(doc, "Developed fraud detection predictive models achieving 94.2% ROC-AUC score across 50M+ financial transactions.")
    add_bullet(doc, "Wrote optimized SQL transformations in Snowflake to extract, clean, and engineer multi-terabyte features.")

    add_section_heading(doc, "Education")
    p_edu = doc.add_paragraph()
    r_edu = p_edu.add_run("Master of Science in Data Science | University of California, Berkeley")
    set_run_font(r_edu, font_size=10)

    out_file = CLIENTS_DIR / "tanuja_resume.docx"
    doc.save(out_file)
    print(f"Created Tanuja resume: {out_file}")


def build_jane_doe_resume():
    doc = Document()
    create_header(
        doc,
        name="JANE DOE",
        contact_line="Dallas, TX | 555-019-2834 | janedoe@example.com | linkedin.com/in/janedoe | Work Auth: H1B / C2C"
    )

    add_section_heading(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    r_sum = p_sum.add_run(
        "Dynamic and detail-oriented Frontend / React Developer with 5+ years of experience building modern, responsive, "
        "and high-performance web applications. Specialized in React 18, TypeScript, Next.js, Redux Toolkit, and Tailwind CSS."
    )
    set_run_font(r_sum, font_size=10.5)

    add_section_heading(doc, "Technical Skills")
    table = doc.add_table(rows=3, cols=2)
    skills = [
        ("Frontend & Frameworks:", "React.js, Next.js, TypeScript, JavaScript (ES6+), HTML5, CSS3, Tailwind CSS, Bootstrap"),
        ("State & APIs:", "Redux Toolkit, Context API, React Query, RESTful APIs, GraphQL, Axios"),
        ("Tools & Build:", "Webpack, Vite, Jest, React Testing Library, Git, GitHub, Docker, Figma")
    ]
    for i, (cat, val) in enumerate(skills):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        r0 = c0.paragraphs[0].add_run(cat)
        set_run_font(r0, bold=True, font_size=10, color_rgb=(24, 43, 73))
        r1 = c1.paragraphs[0].add_run(val)
        set_run_font(r1, font_size=10)

    add_section_heading(doc, "Professional Experience")
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(8)
    r1 = p1.add_run("Senior React Developer | CloudWave Solutions, Dallas, TX (2022 – Present)")
    set_run_font(r1, bold=True, font_size=11, color_rgb=(24, 43, 73))
    add_bullet(doc, "Developed scalable customer-facing SaaS dashboards using React 18, TypeScript, and Tailwind CSS.")
    add_bullet(doc, "Engineered state management with Redux Toolkit and React Query, improving client-side caching.")
    add_bullet(doc, "Implemented unit and integration tests with Jest and React Testing Library maintaining 88%+ code coverage.")

    add_section_heading(doc, "Education")
    p_edu = doc.add_paragraph()
    r_edu = p_edu.add_run("Bachelor of Science in Information Technology | University of North Texas")
    set_run_font(r_edu, font_size=10)

    out_file = CLIENTS_DIR / "jane_doe_resume.docx"
    doc.save(out_file)
    print(f"Created Jane Doe resume: {out_file}")


if __name__ == "__main__":
    CLIENTS_DIR.mkdir(parents=True, exist_ok=True)
    build_safi_resume()
    build_tanuja_resume()
    build_jane_doe_resume()
